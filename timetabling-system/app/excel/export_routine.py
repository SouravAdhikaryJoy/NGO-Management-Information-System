"""Excel export of the solved routine (design doc section 5).

MasterTimetable layout (parseable back by the manual-edit re-import):

    Day: MON
    Room | 1 | 2 | ...        <- slot indexes
    R101 | CS101 (LECTURE) / CSE1A / T01 | ...
    <blank row between day blocks>
"""

from __future__ import annotations

from typing import Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session as DbSession

from app.models import ClassGroup, Course, Room, Session, Teacher, TimeSlot
from app.solver.constraints.registry import hard_violation_report, soft_penalty_breakdown
from app.solver.domain import Problem, Timetable

HEADER_FONT = Font(bold=True)
DAY_ORDER = ["SAT", "SUN", "MON", "TUE", "WED", "THU", "FRI"]


def _cell_text(course_code, session_type, group_code, teacher_code):
    return f"{course_code} ({session_type}) / {group_code} / {teacher_code}"


def _sheet_title(prefix: str, code: str) -> str:
    title = f"{prefix}_{code}"
    for bad in "[]:*?/\\":
        title = title.replace(bad, "-")
    return title[:31]


def export_routine(db: DbSession, problem: Problem, timetable: Timetable, path_or_buffer) -> None:
    rooms = {r.id: r for r in db.query(Room).all()}
    teachers = {t.id: t for t in db.query(Teacher).all()}
    groups = {g.id: g for g in db.query(ClassGroup).all()}
    courses = {c.id: c for c in db.query(Course).all()}
    sessions = {s.id: s for s in db.query(Session).all()}

    days = sorted(
        {t.day_of_week for t in db.query(TimeSlot).all()},
        key=lambda d: DAY_ORDER.index(d) if d in DAY_ORDER else 99,
    )
    slot_indexes = sorted({t.slot_index for t in db.query(TimeSlot).all()})

    # (day, slot_index, room_id) -> session_id, expanded over durations
    occupancy: Dict[tuple, int] = {}
    for sid, placement in timetable.placements.items():
        duration = problem.sessions[sid].duration
        for i in range(placement.index, placement.index + duration):
            occupancy[(placement.day, i, placement.room_id)] = sid

    wb = Workbook()

    # --- MasterTimetable ---
    ws = wb.active
    ws.title = "MasterTimetable"
    row_cursor = 1
    room_list = sorted(rooms.values(), key=lambda r: r.code)
    for day in days:
        ws.cell(row=row_cursor, column=1, value=f"Day: {day}").font = HEADER_FONT
        row_cursor += 1
        ws.cell(row=row_cursor, column=1, value="Room").font = HEADER_FONT
        for j, index in enumerate(slot_indexes, start=2):
            ws.cell(row=row_cursor, column=j, value=index).font = HEADER_FONT
        row_cursor += 1
        for room in room_list:
            ws.cell(row=row_cursor, column=1, value=room.code)
            for j, index in enumerate(slot_indexes, start=2):
                sid = occupancy.get((day, index, room.id))
                if sid is not None:
                    s = sessions[sid]
                    ws.cell(row=row_cursor, column=j, value=_cell_text(
                        courses[s.course_id].code, s.session_type,
                        groups[s.class_group_id].code, teachers[s.teacher_id].code,
                    ))
            row_cursor += 1
        row_cursor += 1  # blank row between day blocks
    ws.column_dimensions["A"].width = 12

    # --- per-teacher and per-group sheets ---
    def personal_sheet(title, owner_filter, describe):
        sheet = wb.create_sheet(title)
        sheet.cell(row=1, column=1, value="Day").font = HEADER_FONT
        for j, index in enumerate(slot_indexes, start=2):
            sheet.cell(row=1, column=j, value=index).font = HEADER_FONT
        for i, day in enumerate(days, start=2):
            sheet.cell(row=i, column=1, value=day)
            for j, index in enumerate(slot_indexes, start=2):
                for (d, idx, room_id), sid in occupancy.items():
                    if d == day and idx == index and owner_filter(sessions[sid]):
                        sheet.cell(row=i, column=j, value=describe(sessions[sid], room_id))
                        break
        sheet.column_dimensions["A"].width = 8

    for teacher in sorted(teachers.values(), key=lambda t: t.code):
        personal_sheet(
            _sheet_title("T", teacher.code),
            lambda s, tid=teacher.id: s.teacher_id == tid,
            lambda s, room_id: (
                f"{courses[s.course_id].code} ({s.session_type}) / "
                f"{groups[s.class_group_id].code} @ {rooms[room_id].code}"
            ),
        )
    for group in sorted(groups.values(), key=lambda g: g.code):
        personal_sheet(
            _sheet_title("G", group.code),
            lambda s, gid=group.id: s.class_group_id == gid,
            lambda s, room_id: (
                f"{courses[s.course_id].code} ({s.session_type}) / "
                f"{teachers[s.teacher_id].code} @ {rooms[room_id].code}"
            ),
        )

    # --- FeasibilityReport ---
    ws = wb.create_sheet("FeasibilityReport")
    for j, header in enumerate(["constraint_key", "description", "violations"], start=1):
        ws.cell(row=1, column=j, value=header).font = HEADER_FONT
    for i, row in enumerate(hard_violation_report(problem, timetable), start=2):
        ws.cell(row=i, column=1, value=row["key"])
        ws.cell(row=i, column=2, value=row["description"])
        ws.cell(row=i, column=3, value=row["violations"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70

    # --- SoftConstraintScoreReport ---
    ws = wb.create_sheet("SoftConstraintScoreReport")
    headers = ["constraint_key", "tier", "weight", "enabled", "raw_penalty", "weighted_penalty"]
    for j, header in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=header).font = HEADER_FONT
    breakdown = soft_penalty_breakdown(problem, timetable)
    field_map = {"constraint_key": "key"}
    i = 2
    for row in breakdown:
        for j, field in enumerate(headers, start=1):
            ws.cell(row=i, column=j, value=row[field_map.get(field, field)])
        i += 1
    tier_totals: Dict[int, float] = {}
    for row in breakdown:
        tier_totals[row["tier"]] = tier_totals.get(row["tier"], 0.0) + row["weighted_penalty"]
    i += 1
    for tier in sorted(tier_totals):
        ws.cell(row=i, column=1, value=f"Tier {tier} subtotal").font = HEADER_FONT
        ws.cell(row=i, column=6, value=tier_totals[tier])
        i += 1
    ws.cell(row=i, column=1, value="TOTAL").font = HEADER_FONT
    ws.cell(row=i, column=6, value=sum(tier_totals.values()))
    ws.column_dimensions["A"].width = 30

    wb.save(path_or_buffer)
