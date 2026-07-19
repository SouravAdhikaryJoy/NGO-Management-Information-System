"""Excel import validation, export round-trip, and manual-edit re-import."""

import io
import random

import pandas as pd
from openpyxl import load_workbook

from app.excel.import_templates import SHEET_ORDER, import_workbook, write_blank_template
from app.excel.export_routine import export_routine
from app.excel.reimport import reimport_master_timetable
from app.excel.sample_data import build_sample_frames, write_sample_workbook
from app.models import ClassGroup, Room, Session, Teacher
from app.solver.constraints.registry import check_all_hard
from app.solver.domain import Placement
from app.solver.loader import (
    generate_sessions,
    load_current_timetable,
    load_problem,
    persist_timetable,
)
from app.solver.phase1_construction import run_phase1


def sample_buffer(scale=1, mutate=None):
    frames = build_sample_frames(scale)
    if mutate:
        mutate(frames)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet, df in frames.items():
            df.to_excel(writer, sheet_name=sheet, index=False)
    buffer.seek(0)
    return buffer


def test_blank_template_has_all_sheets():
    buffer = io.BytesIO()
    write_blank_template(buffer)
    buffer.seek(0)
    wb = load_workbook(buffer)
    assert set(wb.sheetnames) == set(SHEET_ORDER)


def test_import_sample_workbook(db):
    result = import_workbook(db, sample_buffer())
    assert result.ok, result.errors
    assert result.counts["Courses"] == 10
    assert result.counts["ClassGroups"] == 4
    assert db.query(Room).count() == 6
    assert db.query(Teacher).count() == 8


def test_reimport_is_idempotent_upsert(db):
    assert import_workbook(db, sample_buffer()).ok
    teachers_before = db.query(Teacher).count()

    def rename_teacher(frames):
        frames["Teachers"].loc[0, "name"] = "Renamed Teacher"

    result = import_workbook(db, sample_buffer(mutate=rename_teacher))
    assert result.ok
    assert db.query(Teacher).count() == teachers_before  # no duplicates
    renamed = db.query(Teacher).filter(Teacher.code == "D00T01").one()
    assert renamed.name == "Renamed Teacher"


def test_import_reports_row_level_errors_and_writes_nothing(db):
    def corrupt(frames):
        frames["Rooms"]["capacity"] = frames["Rooms"]["capacity"].astype(object)
        frames["Rooms"].loc[0, "capacity"] = "not-a-number"
        frames["Courses"].loc[1, "department_code"] = "NOPE"
        frames["Teachers"].loc[2, "name"] = None

    result = import_workbook(db, sample_buffer(mutate=corrupt))
    assert not result.ok
    described = {(e["sheet"], e["field"]) for e in result.errors}
    assert ("Rooms", "capacity") in described
    assert ("Courses", "department_code") in described
    assert ("Teachers", "name") in described
    # row numbers are Excel-style (header on row 1)
    assert all(e["row"] >= 2 for e in result.errors)
    # all-or-nothing: nothing was committed
    assert db.query(Room).count() == 0
    assert db.query(Teacher).count() == 0


def test_import_missing_required_sheet(db):
    def drop(frames):
        del frames["Rooms"]

    result = import_workbook(db, sample_buffer(mutate=drop))
    assert not result.ok
    assert any(e["sheet"] == "Rooms" and e["field"] == "<sheet>" for e in result.errors)


def solved_state(db):
    assert import_workbook(db, sample_buffer()).ok
    generate_sessions(db)
    problem = load_problem(db)
    p1 = run_phase1(problem, random.Random(42))
    assert p1.feasible
    persist_timetable(db, p1.timetable)
    return problem, p1.timetable


def test_export_workbook_structure(db):
    problem, timetable = solved_state(db)
    buffer = io.BytesIO()
    export_routine(db, problem, timetable, buffer)
    buffer.seek(0)
    wb = load_workbook(buffer)
    names = wb.sheetnames
    assert "MasterTimetable" in names
    assert "FeasibilityReport" in names
    assert "SoftConstraintScoreReport" in names
    teacher_sheets = [n for n in names if n.startswith("T_")]
    group_sheets = [n for n in names if n.startswith("G_")]
    assert len(teacher_sheets) == db.query(Teacher).count()
    assert len(group_sheets) == db.query(ClassGroup).count()
    # feasibility report shows 14 hard constraints, all zero
    ws = wb["FeasibilityReport"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 14
    assert all(row[2] == 0 for row in rows)


def _find_feasible_edit(db, problem, timetable):
    """Pick a movable session and a feasible alternative placement."""
    for sid, current in timetable.placements.items():
        session = problem.sessions[sid]
        if session.duration != 1:
            continue
        for placement in problem.all_placements(session):
            if placement == current:
                continue
            timetable.remove(sid)
            ok = check_all_hard(session, placement, problem, timetable)
            timetable.place(sid, current)
            if ok:
                return sid, current, placement
    raise AssertionError("no feasible edit found")


def _move_cell(wb, problem, db, sid, old, new):
    """Rewrite the MasterTimetable grid to move session `sid` old->new."""
    ws = wb["MasterTimetable"]
    rooms = {r.id: r.code for r in db.query(Room).all()}
    day = None
    header_cols = {}
    old_text = None
    for row in ws.iter_rows():
        first = row[0].value
        if isinstance(first, str) and first.startswith("Day: "):
            day = first[5:]
            continue
        if first == "Room":
            header_cols = {cell.column: cell.value for cell in row[1:] if cell.value}
            continue
        if first is None:
            continue
        for cell in row[1:]:
            if day == old.day and header_cols.get(cell.column) == old.index \
                    and str(first) == rooms[old.room_id] and cell.value:
                old_text = cell.value
                cell.value = None
    assert old_text is not None
    day = None
    for row in ws.iter_rows():
        first = row[0].value
        if isinstance(first, str) and first.startswith("Day: "):
            day = first[5:]
            continue
        if first == "Room":
            header_cols = {cell.column: cell.value for cell in row[1:] if cell.value}
            continue
        if first is None:
            continue
        for cell in row[1:]:
            if day == new.day and header_cols.get(cell.column) == new.index \
                    and str(first) == rooms[new.room_id]:
                cell.value = old_text


def test_manual_edit_reimport_applies_valid_edit(db):
    problem, timetable = solved_state(db)
    sid, old, new = _find_feasible_edit(db, problem, timetable)
    buffer = io.BytesIO()
    export_routine(db, problem, timetable, buffer)
    buffer.seek(0)
    wb = load_workbook(buffer)
    _move_cell(wb, problem, db, sid, old, new)
    edited = io.BytesIO()
    wb.save(edited)
    edited.seek(0)

    result = reimport_master_timetable(db, edited)
    assert result["rejected"] == [], result["rejected"]
    assert any(a["session_id"] == sid for a in result["applied"])
    session = db.get(Session, sid)
    assert (session.day_of_week, session.slot_index, session.room_id) == (
        new.day, new.index, new.room_id,
    )
    assert result["soft_report"] is not None


def test_manual_edit_reimport_rejects_hard_violation(db):
    problem, timetable = solved_state(db)
    # move a session onto a slot where its own teacher already teaches
    sid, violating = None, None
    for candidate, current in timetable.placements.items():
        session = problem.sessions[candidate]
        if session.duration != 1:
            continue
        teacher_days = timetable.teacher_slots.get(session.teacher_id, {})
        for day, indexes in teacher_days.items():
            for index in indexes:
                if (day, index) != (current.day, current.index):
                    for room_id in problem.rooms:
                        occupant = timetable.room_busy.get((room_id, day, index))
                        room_ok = occupant is None
                        if room_ok:
                            sid, violating = candidate, Placement(day, index, room_id)
                            break
                if violating:
                    break
            if violating:
                break
        if violating:
            break
    assert violating is not None

    buffer = io.BytesIO()
    export_routine(db, problem, timetable, buffer)
    buffer.seek(0)
    wb = load_workbook(buffer)
    old = timetable.placements[sid]
    _move_cell(wb, problem, db, sid, old, violating)
    edited = io.BytesIO()
    wb.save(edited)
    edited.seek(0)

    result = reimport_master_timetable(db, edited)
    assert any("H02_teacher_clash" in str(r.get("error")) for r in result["rejected"]), result
    # the rejected session kept its original assignment
    session = db.get(Session, sid)
    assert (session.day_of_week, session.slot_index) == (old.day, old.index)
